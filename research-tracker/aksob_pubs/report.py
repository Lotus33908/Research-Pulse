"""Excel workbook + CSV export."""
from __future__ import annotations

import csv
import json
import logging
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .config import Config
from .db import Database
from .summaries import (QUARTILES, coverage_status, department_summary, faculty_summary, is_covered, journal_summary,
                        on_scholar, pub_key, rank_status, unique_pubs, year_summary)
from .utils import fmt_issn

log = logging.getLogger("aksob")

HEADER_FILL = PatternFill("solid", fgColor="184F95")
HEADER_FONT = Font(bold=True, color="FFFFFF")
Q_FILL = {
    "Q1": (PatternFill("solid", fgColor="184F95"), Font(color="FFFFFF", bold=True)),
    "Q2": (PatternFill("solid", fgColor="2A78D6"), Font(color="FFFFFF", bold=True)),
    "Q3": (PatternFill("solid", fgColor="86B6EF"), Font(color="0B0B0B")),
    "Q4": (PatternFill("solid", fgColor="CDE2FB"), Font(color="0B0B0B")),
}


def write_table(ws: Worksheet, headers: list[str], rows: Iterable[Iterable[Any]], widths: dict[str, int] | None = None,
                quartile_col: str | None = None, wrap_cols: tuple[str, ...] = (), link_col: str | None = None) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    n = 0
    for row in rows:
        ws.append(list(row))
        n += 1
    ws.freeze_panes = "A2"
    if n:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n + 1}"
    widths = widths or {}
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, max(10, min(28, len(h) + 4)))
    if quartile_col and quartile_col in headers:
        ci = headers.index(quartile_col) + 1
        for r in range(2, n + 2):
            cell = ws.cell(row=r, column=ci)
            style = Q_FILL.get(str(cell.value))
            if style:
                cell.fill, cell.font = style
                cell.alignment = Alignment(horizontal="center")
    for col in wrap_cols:
        if col in headers:
            ci = headers.index(col) + 1
            for r in range(2, n + 2):
                ws.cell(row=r, column=ci).alignment = Alignment(wrap_text=True, vertical="top")
    if link_col and link_col in headers:
        ci = headers.index(link_col) + 1
        for r in range(2, n + 2):
            cell = ws.cell(row=r, column=ci)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.font = Font(color="1C5CAB", underline="single")


def pct(x: float | None) -> str:
    return f"{x:.0%}" if x is not None else ""


def build_report(cfg: Config, db: Database, scimago_years: list[int] | None = None) -> Path:
    if scimago_years is None:
        from .scimago import cached_years
        scimago_years = cached_years(cfg.scimago_dir)
    cfg.output_dir.mkdir(parents=True, exist_ok=True)
    faculty = db.faculty()
    pubs = db.publications_joined()
    fac_rows = faculty_summary(faculty, pubs)
    dept_rows = department_summary(faculty, pubs)
    year_rows = year_summary(pubs)
    jr_rows = journal_summary(pubs)
    ups = unique_pubs(pubs)
    generated = datetime.now().strftime("%Y-%m-%d %H:%M")

    wb = Workbook()
    # ------------------------------------------------------------- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{cfg.school_name} - Research Publications"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generated {generated}. Publications from {cfg.year_from} onwards. Quartiles: Scimago Journal Rank (best quartile, year of publication)."
    c = Counter(rank_status(p) for p in ups)
    imports = {f["slug"] for f in faculty if f["scholar_import_at"]}
    kpis = [
        ("Faculty included", len(faculty)),
        ("Faculty covered (OpenAlex record, ORCID or Scholar export)", sum(1 for f in faculty if is_covered(f))),
        ("Faculty NOT covered - see Needs Review", sum(1 for f in faculty if not is_covered(f))),
        ("Faculty with a Google Scholar export imported", len(imports)),
        ("Faculty with publications", len({p['faculty_slug'] for p in pubs})),
        ("Unique publications", len(ups)),
        ("  of which in a Scholar export", sum(1 for p in ups if on_scholar(p))),
        ("  found by the indexes but missing from the person's Scholar export",
         sum(1 for p in pubs if p["faculty_slug"] in imports and not on_scholar(p))),
        ("Q1", c["Q1"]), ("Q2", c["Q2"]), ("Q3", c["Q3"]), ("Q4", c["Q4"]),
        ("Journal items not matched to Scimago", c["Unranked"] + c["Indexed, no quartile"]),
        ("Non-journal items (books, working papers, proceedings...)", c["Not applicable"]),
        ("Q1+Q2 share of journal items", pct((c["Q1"] + c["Q2"]) / (sum(c[q] for q in QUARTILES) + c["Unranked"] + c["Indexed, no quartile"]) if pubs else None)),
    ]
    for i, (k, v) in enumerate(kpis, start=4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 14
    r0 = 4 + len(kpis) + 1
    ws.cell(row=r0, column=1, value="By department (unique publications)").font = Font(bold=True, size=12)
    hdr = ["Department", "Faculty", "Publications", "Q1", "Q2", "Q3", "Q4", "Unranked", "Non-journal", "Q1+Q2", "Faculty with Q1/Q2 in last 3 yrs"]
    for j, h in enumerate(hdr, 1):
        cell = ws.cell(row=r0 + 1, column=j, value=h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for i, d in enumerate(dept_rows, start=r0 + 2):
        vals = [d["department"], d["faculty"], d["total"], d["Q1"], d["Q2"], d["Q3"], d["Q4"], d["unranked"], d["na"], d["q12"], d["faculty_q12_recent"]]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
    r1 = r0 + 2 + len(dept_rows) + 1
    ws.cell(row=r1, column=1, value="By year (unique publications)").font = Font(bold=True, size=12)
    hdr = ["Year", "Publications", "Q1", "Q2", "Q3", "Q4", "Unranked", "Non-journal"]
    for j, h in enumerate(hdr, 1):
        cell = ws.cell(row=r1 + 1, column=j, value=h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for i, y in enumerate(year_rows, start=r1 + 2):
        for j, v in enumerate([y["year"], y["total"], y["Q1"], y["Q2"], y["Q3"], y["Q4"], y["unranked"], y["na"]], 1):
            ws.cell(row=i, column=j, value=v)
    for col in "CDEFGHIJK":
        ws.column_dimensions[col].width = 13

    # ------------------------------------------------------------ Faculty
    ws = wb.create_sheet("Faculty Summary")
    headers = ["Faculty", "Department", "Category", "Title", "Publications", "Q1", "Q2", "Q3", "Q4", "Unranked",
               "Non-journal", "Q1+Q2", "Q1+Q2 share", "Q1+Q2 last 5 yrs", "ABDC A*/A", "In Scholar export",
               "Missing from Scholar export", "Coverage", "Citations", "h-index", "Metrics source", "OpenAlex",
               "ORCID", "Google Scholar", "Fetch status"]
    write_table(ws, headers, (
        [r["name"], r["department"], r["category"], r["title"], r["total"], r["Q1"], r["Q2"], r["Q3"], r["Q4"],
         r["unranked"], r["na"], r["q12"], pct(r["q12_share"]), r["q12_recent"], r["abdc_top"],
         r["on_scholar"] if r["scholar_articles"] is not None else "", r["not_in_scholar_export"] if r["not_in_scholar_export"] is not None else "",
         r["coverage"], r["citations"], r["h_index"], r["metrics_source"],
         f"https://openalex.org/authors/{r['openalex_id']}" if r["openalex_id"] else "",
         f"https://orcid.org/{r['orcid']}" if r["orcid"] else "",
         f"https://scholar.google.com/citations?user={r['scholar_id']}" if r["scholar_id"] else "", r["fetch_status"] or ""]
        for r in fac_rows),
        widths={"Faculty": 26, "Department": 26, "Title": 34, "Coverage": 44, "OpenAlex": 36, "ORCID": 30,
                "Google Scholar": 44, "Fetch status": 44},
        link_col="OpenAlex")

    # ------------------------------------------------------- Publications
    ws = wb.create_sheet("Publications")
    headers = ["Faculty", "Department", "Year", "Title", "Journal / Venue", "Quartile", "SJR", "ABDC", "Type",
               "Citations", "In Scholar export", "DOI", "ISSN", "Scimago categories", "Scimago year", "Match",
               "Sources", "Venue as listed"]
    write_table(ws, headers, (
        [p["faculty_name"], p["department"], p["year"], p["title"], p["scimago_title"] or p["venue"] or p["venue_raw"],
         p["quartile"] or ("n/a" if (p["match_method"] or "") == "n/a" else ""), p["sjr"], p["abdc"] or "",
         p["type"] or "", (p["cited_by_scholar"] if p["cited_by_scholar"] is not None else p["cited_by_openalex"]),
         "Yes" if on_scholar(p) else "No",
         f"https://doi.org/{p['doi']}" if p["doi"] else "",
         ", ".join(fmt_issn(i) for i in json.loads(p["issns"])) if p["issns"] else "",
         p["categories"] or "", p["scimago_year"], p["match_method"] or "", p["sources"] or "", p["venue_raw"] or ""]
        for p in pubs),
        widths={"Faculty": 22, "Department": 22, "Title": 60, "Journal / Venue": 38, "DOI": 36,
                "Scimago categories": 40, "Venue as listed": 36, "Year": 7, "Quartile": 9},
        quartile_col="Quartile", wrap_cols=("Title",), link_col="DOI")

    # ----------------------------------------------------------- Journals
    ws = wb.create_sheet("Journals")
    headers = ["Journal", "Papers", "Quartile", "Scimago year", "SJR", "ABDC", "Scimago categories", "Publisher", "AKSOB authors"]
    write_table(ws, headers, (
        [j["journal"], j["papers"], j["quartile"], j["scimago_year"], j["sjr"], j["abdc"], j["categories"], j["publisher"], j["faculty"]]
        for j in jr_rows),
        widths={"Journal": 46, "Scimago categories": 44, "Publisher": 28, "AKSOB authors": 40},
        quartile_col="Quartile")

    # ------------------------------------------------------- Needs review
    ws = wb.create_sheet("Needs Review")
    review_rows: list[list[Any]] = []
    for f in faculty:                                   # coverage gaps come first - they matter most
        if not is_covered(f):
            review_rows.append([f"COVERAGE GAP - {coverage_status(f)}. Add openalex_id or orcid in data/faculty.csv, "
                                "or drop the person's Scholar BibTeX export in data/scholar_exports/",
                                f["name"], "", "", "", "", "", ""])
    for p in pubs:
        if p["faculty_slug"] in imports and not on_scholar(p):
            review_rows.append(["Found by OpenAlex/ORCID but missing from the person's Google Scholar export - ask them to add it, or ignore",
                                p["faculty_name"], p["year"], p["title"], p["venue"] or p["venue_raw"], p["scimago_title"],
                                p["quartile"], ""])
    for p in pubs:
        m = p["match_method"] or ""
        if m.endswith("fuzzy"):
            review_rows.append(["Fuzzy journal match - confirm", p["faculty_name"], p["year"], p["title"], p["venue_raw"] or p["venue"],
                                p["scimago_title"], p["quartile"], p["match_score"]])
    unranked: dict[str, list] = defaultdict(list)
    for p in pubs:
        if rank_status(p) == "Unranked":
            unranked[(p["venue"] or p["venue_raw"] or "(no venue)")].append(p)
    for venue, ps in sorted(unranked.items(), key=lambda kv: -len(kv[1])):
        review_rows.append([f"Venue not found in Scimago ({len(ps)} papers) - add to journal_overrides.csv if it is a journal",
                            ", ".join(sorted({p['faculty_name'] for p in ps}))[:120], ps[0]["year"], ps[0]["title"], venue, "", "", ""])
    for f in faculty:
        if f["fetch_status"] and "error" in f["fetch_status"]:
            review_rows.append(["Faculty record - last fetch reported an error: " + f["fetch_status"], f["name"], "", "", "", "", "", ""])
    write_table(ws, ["Issue", "Faculty", "Year", "Title", "Venue as listed", "Matched Scimago title", "Quartile", "Score"],
                review_rows, widths={"Issue": 70, "Faculty": 26, "Title": 56, "Venue as listed": 36, "Matched Scimago title": 36})

    # ------------------------------------------------------------- Roster
    ws = wb.create_sheet("Roster")
    write_table(ws, ["Name", "Department", "Category", "Title", "OpenAlex id", "ORCID", "Google Scholar id", "Scholar export", "Profile"],
                ([f["name"], f["department"], f["category"], f["title"], f["openalex_id"], f["orcid"], f["scholar_id"],
                  (str(f["scholar_import_at"])[:10] if f["scholar_import_at"] else ""), f["profile_url"]]
                 for f in faculty), widths={"Name": 26, "Department": 26, "Title": 36, "Profile": 60})

    # -------------------------------------------------------------- About
    ws = wb.create_sheet("About")
    lines = [
        "AKSOB Research Tracker",
        f"Generated: {generated}",
        f"Publication years: {cfg.year_from} onwards",
        f"Sources: {', '.join(cfg.sources)}",
        f"Scimago years loaded: {', '.join(map(str, scimago_years)) if scimago_years else 'see run log'}",
        "",
        "How the list is built and quartiles assigned:",
        "1. Each faculty member's papers come from their OpenAlex author record (matched by name + LAU affiliation) and their ORCID record when known.",
        "   Optionally, a Google Scholar BibTeX export dropped in data/scholar_exports/ is imported too and used to check nothing is missing.",
        "2. Titles without a DOI/ISSN are resolved via OpenAlex, then Crossref (journal version preferred over preprints).",
        "3. The journal's ISSN is looked up in the Scimago Journal Rank list of the paper's publication year (nearest available year otherwise).",
        "4. If no ISSN is known, the venue name is matched exactly, then with a strict fuzzy match (flagged in 'Needs Review').",
        "5. 'Quartile' is Scimago's SJR Best Quartile across the journal's subject categories. ABDC is the Australian Business Deans Council rating (via OpenAlex).",
        "6. 'n/a' marks books, chapters, working papers, proceedings and other non-journal items. Blank = journal-like venue not found in Scimago.",
        "Manual fixes: data/journal_overrides.csv (venue -> ISSN or quartile), data/faculty.csv (ids, include flag).",
    ]
    for i, line in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=line).font = Font(bold=(i == 1 or line.endswith(":")))
    ws.column_dimensions["A"].width = 140

    out = cfg.output_dir / "AKSOB_Publications.xlsx"
    wb.save(out)
    log.info("Excel report written to %s", out)
    write_csv(cfg.output_dir / "publications.csv", pubs)
    return out


def write_csv(path: Path, pubs: list) -> None:
    cols = ["faculty_name", "department", "year", "title", "venue", "venue_raw", "scimago_title", "quartile", "sjr", "abdc",
            "type", "cited_by_scholar", "cited_by_openalex", "doi", "issns", "categories", "scimago_year", "match_method",
            "match_score", "sources", "authors", "pub_key"]
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(cols)
        for p in pubs:
            row = []
            for c in cols:
                if c == "pub_key":
                    row.append(pub_key(p))
                elif c == "issns":
                    row.append(", ".join(fmt_issn(i) for i in json.loads(p["issns"])) if p["issns"] else "")
                else:
                    row.append(p[c])
            w.writerow(row)
    log.info("CSV written to %s", path)
